#!/usr/bin/env python3
"""结构指标基线：直接从 .pptx 里量"这册排得乱不乱"。

用途：在没有桌面桥（生成链路）的情况下，先用真实产物校准 `craft.ts` 的阈值，
特别是留白率区间、主角面积区间、对齐轴数与行长度。

跑法：
  /Users/xjwang/.dsh/dsh-runtimes/dsh-primary-runtime/dependencies/python/bin/python3 \
      apps/gateway/scripts/slides-structure-baseline.py <deck.pptx> [更多.pptx ...]

口径（与 apps/gateway/src/modules/slides-layout/page-structure.ts 对齐）：
  - 画布按 1280×720 归一，1pt = 4/3 px；
  - 全屏色块（宽高都 ≥98%）算背景，不算内容块；
  - 主角面积 = 最大内容块面积 / 画布面积；
  - 留白率 = 1 − 内容块并集面积 / 画布面积；
  - 对齐轴 = 文字块左边缘按 12px 容差聚类后的条数；
  - measure = 最长一行中文字数（字宽按 1.35 × 字号估算，与渲染器同一系数）；
  - 字号档 = 相差 <12% 视为同档。
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

CANVAS_W = 1280.0
CANVAS_H = 720.0
EMU_PER_INCH = 914400.0


def union_area(rects, box):
    """矩形并集面积（坐标压缩）。"""
    clipped = []
    for x, y, w, h in rects:
        x0, y0 = max(x, box[0]), max(y, box[1])
        x1, y1 = min(x + w, box[0] + box[2]), min(y + h, box[1] + box[3])
        if x1 > x0 and y1 > y0:
            clipped.append((x0, y0, x1, y1))
    if not clipped:
        return 0.0
    xs = sorted({value for rect in clipped for value in (rect[0], rect[2])})
    ys = sorted({value for rect in clipped for value in (rect[1], rect[3])})
    area = 0.0
    for i in range(len(xs) - 1):
        for j in range(len(ys) - 1):
            cx, cy = xs[i], ys[j]
            cw, ch = xs[i + 1] - cx, ys[j + 1] - cy
            if cw <= 0 or ch <= 0:
                continue
            for x0, y0, x1, y1 in clipped:
                if x0 <= cx and y0 <= cy and x1 >= cx + cw and y1 >= cy + ch:
                    area += cw * ch
                    break
    return area


def shape_font_sizes(shape):
    sizes = []
    if not shape.has_text_frame:
        return sizes
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
    return sizes


def shape_text(shape):
    if not shape.has_text_frame:
        return ""
    return shape.text_frame.text.strip()


def measure_deck(path: str):
    prs = Presentation(path)
    slide_w = float(prs.slide_width)
    slide_h = float(prs.slide_height)
    # EMU → px（按 1280 宽归一）
    scale = CANVAS_W / slide_w
    box = (0.0, 0.0, CANVAS_W, slide_h * scale)

    rows = []
    for index, slide in enumerate(prs.slides, start=1):
        blocks = []
        text_blocks = []
        sizes = []
        measure_max = 0
        for shape in slide.shapes:
            if shape.width is None or shape.height is None:
                continue
            x = float(shape.left) * scale
            y = float(shape.top) * scale
            w = float(shape.width) * scale
            h = float(shape.height) * scale
            if w <= 0 or h <= 0:
                continue
            full_bleed = w >= box[2] * 0.98 and h >= box[3] * 0.98
            if full_bleed:
                continue
            blocks.append((x, y, w, h))
            font_sizes = shape_font_sizes(shape)
            if font_sizes:
                sizes.extend(font_sizes)
                text_blocks.append((x, shape_text(shape), max(font_sizes), w))
        canvas_area = box[2] * box[3]
        hero = max((w * h for _x, _y, w, h in blocks), default=0.0) / canvas_area
        whitespace = 1.0 - union_area(blocks, box) / canvas_area

        axes = []
        for x, _text, _size, _w in sorted(text_blocks):
            if not axes or x - axes[-1] > 12:
                axes.append(x)

        for _x, text, size, width in text_blocks:
            if not text or size <= 0:
                continue
            char_w = 1.35 * size * (1280.0 / 960.0)  # pt → px
            lines = max(1, -(-int(len(text) * char_w) // max(1, int(width))))
            measure_max = max(measure_max, round(len(text) / lines))

        tiers = []
        for size in sorted(set(sizes), reverse=True):
            if not tiers or tiers[-1] / size > 1.12:
                tiers.append(size)

        rows.append({
            "page": index,
            "blocks": len(blocks),
            "hero": hero,
            "whitespace": whitespace,
            "axes": len(axes),
            "measure": measure_max,
            "tiers": len(tiers),
        })
    return rows


def summarize(name: str, rows):
    def avg(key):
        return sum(row[key] for row in rows) / len(rows)

    print(f"\n### {name}（{len(rows)} 页）\n")
    print("| 页 | 内容块 | 主角面积 | 留白率 | 对齐轴 | measure | 字号档 |")
    print("| --- | --- | --- | --- | --- | --- | --- |")
    for row in rows:
        print(
            f"| {row['page']} | {row['blocks']} | {row['hero'] * 100:.0f}% | {row['whitespace'] * 100:.0f}% "
            f"| {row['axes']} | {row['measure']} | {row['tiers']} |"
        )
    print(
        f"\n页均：主角 {avg('hero') * 100:.0f}% · 留白 {avg('whitespace') * 100:.0f}% · "
        f"对齐轴 max {max(r['axes'] for r in rows)} · measure max {max(r['measure'] for r in rows)} · "
        f"字号档 max {max(r['tiers'] for r in rows)}"
    )


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 1
    for path in argv[1:]:
        summarize(Path(path).name, measure_deck(path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
